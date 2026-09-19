"""Lister FB list: opt-in per item, review -> ready -> Facebook's workbook -> listed, and the
template itself (categories, conditions, whole-dollar prices, 50-row limit)."""

import io
import unittest

import openpyxl

import fb_lister_template as tpl
from test_lister import UPC, ListerTestCase

HEADERS = {'X-Sweet-Shelves-Lister': '1'}


class FbTemplateTests(unittest.TestCase):
    def test_category_suggestions_use_title_and_ebay_path(self):
        found = [s['category'] for s in tpl.suggest_categories('Ninja Air Fryer 4 Quart', '')]
        self.assertEqual(found[0], 'Home & Kitchen//Kitchen & Dining//Small Kitchen Appliances//Deep Fryers & Air Fryers')
        found = [s['category'] for s in tpl.suggest_categories('Dewalt 20V Drill', 'Home & Garden > Power Tools > Drills')]
        self.assertEqual(found[0], 'Tools & Home Improvement//Tools//Power Tools//Drills')
        self.assertEqual(len(tpl.categories()), 1870)

    def test_conditions_prices_and_text(self):
        self.assertEqual(tpl.fb_condition('NEW_OTHER'), 'New')
        self.assertEqual(tpl.fb_condition('USED_ACCEPTABLE'), 'Used - Fair')
        self.assertEqual(tpl.fb_condition('Used - Good'), 'Used - Good')
        self.assertEqual((tpl.whole_dollars('24.50'), tpl.whole_dollars('$1,299.6'), tpl.whole_dollars('0')), (24, 1300, None))
        self.assertEqual(tpl.plain_text('<p>A <b>plate</b></p><ul><li>Safe</li></ul>', 100), 'A plate\n• Safe')

    def test_workbook_is_facebooks_template_filled_from_row_5(self):
        rows = [{'upc': str(i), 'title': f'Item {i}', 'price': 10 + i, 'condition': 'New', 'description': 'd',
                 'category': tpl.categories()[0]} for i in range(3)]
        book = openpyxl.load_workbook(io.BytesIO(tpl.build_workbook(rows)))
        sheet = book[tpl.SHEET]
        self.assertEqual([c.value for c in sheet[4]][:5], ['TITLE', 'PRICE', 'CONDITION', 'DESCRIPTION', 'CATEGORY'])
        self.assertEqual([c.value for c in sheet[5]][:3], ['Item 0', 10, 'New'])
        self.assertIsNone(sheet.cell(8, 1).value, "Facebook's example row is gone")
        self.assertIn('VALIDATION', book.sheetnames)
        with self.assertRaises(ValueError):
            tpl.build_workbook(rows * 17)  # 51 rows
        with self.assertRaises(ValueError):
            tpl.build_workbook([dict(rows[0], condition='Broken')])


class ListerFbTests(ListerTestCase):
    def fb(self, path, body=None):
        if body is None:
            return self.client.get('/api/lister/fb' + path)
        return self.client.post('/api/lister/fb' + path, json=body, headers=HEADERS)

    def plus_fb(self, upc=UPC, *, on=True, fresh=False):
        return self.client.post(f'/api/lister/queue/{upc}/store', json={'platform': 'fb', 'on': on, 'only': True, 'fresh': fresh},
                                headers=HEADERS)

    def test_fb_is_opt_in_and_a_fresh_add_stays_off_ebay_and_amazon(self):
        # Before "+ FB": no fb key, eBay/Amazon exactly as before.
        states = self.client.post('/api/lister/queue-stores', json={'upcs': [UPC]}).get_json()['items'][UPC]
        self.assertNotIn('fb', states)
        self.assertEqual(self.fb('/queue').get_json()['items'], [])

        res = self.plus_fb(fresh=True)
        self.assertEqual(res.status_code, 200, res.get_json())
        state = res.get_json()['state']
        self.assertEqual((state['fb'], state['ebay'], state['amazon']), ('on', 'off', 'off'))
        self.assertEqual([i['upc'] for i in self.fb('/queue').get_json()['items']], [UPC])
        self.assertEqual(self.client.get('/api/lister/queue?platform=ebay').get_json()['items'], [])

        # Taking it off again.
        self.assertNotIn('fb', self.plus_fb(on=False).get_json()['state'])
        self.assertEqual(self.fb('/queue').get_json()['items'], [])

    def test_adding_fb_to_an_item_already_queued_keeps_ebay_and_amazon(self):
        state = self.plus_fb(fresh=False).get_json()['state']
        self.assertEqual((state['fb'], state['ebay'], state['amazon']), ('on', 'on', 'on'))

    def test_panel_may_report_fb(self):
        dan = dict(HEADERS, **{'Cf-Access-Authenticated-User-Email': 'dan@example.com'})
        res = self.client.post('/api/lister/panel', json={'platform': 'fb'}, headers=dan)
        self.assertEqual(res.status_code, 200, res.get_json())
        self.assertEqual(self.client.get('/api/lister/panel', headers=dan).get_json()['platform'], 'fb')

    def test_review_ready_build_list(self):
        self.plus_fb(fresh=True)
        item = self.fb(f'/item/{UPC}').get_json()
        draft = item['draft']
        self.assertEqual(draft['title'], 'Lenox Butterfly Meadow Dinner Plate')
        self.assertEqual((draft['price'], draft['condition']), (24, 'New'), 'whole dollars; NEW_OTHER -> New')
        self.assertIn('Butterfly Meadow', draft['description'])
        self.assertNotIn('<', draft['description'])
        self.assertIn(draft['category'], tpl.categories())
        self.assertEqual(item['conditions'], list(tpl.CONDITIONS))

        # Ready refuses a bad row, accepts a good one.
        bad = self.fb(f'/item/{UPC}', dict(draft, price='', ready=True))
        self.assertEqual(bad.status_code, 400)
        ok = self.fb(f'/item/{UPC}', dict(draft, price=25, ready=True))
        self.assertEqual(ok.get_json()['status'], 'ready', ok.get_json())

        built = self.fb('/build', {})
        self.assertEqual(built.status_code, 200, built.get_json())
        batch = built.get_json()['batchId']
        self.assertEqual(self.fb('/queue').get_json()['items'][0]['status'], 'in_template')
        self.assertEqual(self.fb(f'/item/{UPC}', dict(draft, ready=True)).status_code, 409, 'locked while in a workbook')

        workbook = self.client.get(f'/api/lister/fb/batch/{batch}.xlsx')
        self.assertEqual(workbook.status_code, 200)
        sheet = openpyxl.load_workbook(io.BytesIO(workbook.data))[tpl.SHEET]
        self.assertEqual([c.value for c in sheet[5]][:3], ['Lenox Butterfly Meadow Dinner Plate', 25, 'New'])
        # The same rows as the CSV Facebook's multiple-listings page takes (the panel uploads this one).
        csv_file = self.client.get(f'/api/lister/fb/batch/{batch}.csv')
        self.assertEqual(csv_file.status_code, 200)
        self.assertTrue(csv_file.mimetype.startswith('text/csv'))
        lines = csv_file.data.decode('utf-8-sig').split('\r\n')
        self.assertEqual(lines[0], 'TITLE,PRICE,CONDITION,DESCRIPTION,CATEGORY')
        self.assertTrue(lines[1].startswith('Lenox Butterfly Meadow Dinner Plate,25,New,'), lines[1])

        tracked = []
        self.lister.fb_list.track_listing = lambda upc, **kw: tracked.append((upc, kw))
        listed = self.fb(f'/batch/{batch}/listed', {})
        self.assertEqual(listed.get_json()['listed'], 1, listed.get_json())
        self.assertEqual(tracked[0][0], UPC)
        self.assertIn(('facebook', UPC), self.bol_marks)
        self.assertEqual(self.fb('/queue').get_json()['items'], [])
        states = self.client.post('/api/lister/queue-stores', json={'upcs': [UPC]}).get_json()['items'][UPC]
        self.assertEqual(states['fb'], 'listed')
        self.assertEqual(self.fb(f'/batch/{batch}/listed', {}).status_code, 409)

    def test_cancelled_workbook_puts_items_back_to_ready(self):
        self.plus_fb(fresh=True)
        draft = self.fb(f'/item/{UPC}').get_json()['draft']
        self.fb(f'/item/{UPC}', dict(draft, ready=True))
        batch = self.fb('/build', {}).get_json()['batchId']
        self.assertEqual(self.fb(f'/batch/{batch}/cancel', {}).status_code, 200)
        self.assertEqual(self.fb('/queue').get_json()['items'][0]['status'], 'ready')
        self.assertEqual(self.fb('/build', {'upcs': ['000000000000']}).status_code, 400, 'nothing ready among those')

    def test_prices_come_from_ebay_comps_and_amazon_offers(self):
        from flask import jsonify
        self.plus_fb()
        # Nothing wired: both sources say so, the search link still works, and the call never fails.
        data = self.fb(f'/item/{UPC}/prices').get_json()
        self.assertTrue(data['success'])
        self.assertEqual(data['searchUrl'], 'https://www.google.com/search?q=' + UPC)
        by = {s['source']: s for s in data['sources']}
        self.assertIn('not available', by['ebay']['error'])
        self.assertIn('not available', by['amazon']['error'])
        self.assertEqual(self.fb('/item/000000000000/prices').status_code, 404)

        fb = self.lister.fb_list
        comps_calls = []

        def fake_comps():
            from flask import request as flask_request
            comps_calls.append(dict(flask_request.args))
            return jsonify({'success': True, 'results': [
                {'price': {'value': '30.00'}}, {'price': {'value': '18.5'}}, {'price': {'value': 'n/a'}}, {'price': {'value': '24'}}]})

        class FakeProducts:
            def __init__(self, **kw):
                pass

            def get_item_offers(self, asin, ItemCondition='New'):
                assert asin == 'B0TESTASIN'
                return type('R', (), {'payload': {'Summary': {
                    'LowestPrices': [{'condition': 'new', 'LandedPrice': {'Amount': '21.99'}}, {'condition': 'used', 'LandedPrice': {'Amount': '9'}}],
                    'BuyBoxPrices': [{'condition': 'New', 'ListingPrice': {'Amount': '22.49'}}],
                    'NumberOfOffers': [{'condition': 'new', 'OfferCount': 4}]}}})()

        import sys, types
        fake_api = types.ModuleType('sp_api.api'); fake_api.Products = FakeProducts
        fake_pkg = types.ModuleType('sp_api'); fake_pkg.api = fake_api
        saved = {k: sys.modules.get(k) for k in ('sp_api', 'sp_api.api')}
        sys.modules['sp_api'], sys.modules['sp_api.api'] = fake_pkg, fake_api
        fb.ebay_comps_view = fake_comps
        fb.amazon_context = lambda: ({'refresh_token': 'x'}, 'seller', 'ATVPDKIKX0DER', 'US')
        try:
            data = self.fb(f'/item/{UPC}/prices').get_json()
        finally:
            for k, v in saved.items():
                if v is None:
                    sys.modules.pop(k, None)
                else:
                    sys.modules[k] = v
        by = {s['source']: s for s in data['sources']}
        self.assertEqual(comps_calls[0]['upc'], UPC)
        self.assertEqual((by['ebay']['count'], by['ebay']['low'], by['ebay']['median'], by['ebay']['high'], by['ebay']['suggested']),
                         (3, 18.5, 24.0, 30.0, 24))
        self.assertEqual((by['amazon']['asin'], by['amazon']['lowest'], by['amazon']['buyBox'], by['amazon']['count'], by['amazon']['suggested']),
                         ('B0TESTASIN', 21.99, 22.49, 4, 22))
        self.assertEqual(by['amazon']['url'], 'https://www.amazon.com/dp/B0TESTASIN')

        # A dead pricing API reports itself on that one source; the other still answers.
        fb.amazon_context = lambda: (_ for _ in ()).throw(RuntimeError('no creds'))
        by = {s['source']: s for s in self.fb(f'/item/{UPC}/prices').get_json()['sources']}
        self.assertEqual(by['ebay']['suggested'], 24)
        self.assertEqual(by['amazon']['error'], 'no creds')

    def test_category_search_and_guard(self):
        found = self.fb('/categories?q=air fryer').get_json()['categories']
        self.assertIn('Home & Kitchen//Kitchen & Dining//Small Kitchen Appliances//Deep Fryers & Air Fryers', found)
        refused = self.client.post('/api/lister/fb/build', json={}, headers={'Sec-Fetch-Site': 'cross-site'})
        self.assertEqual(refused.status_code, 403)


def load_tests(loader, tests, pattern):
    """Only this file's tests: ListerFbTests inherits the whole test_lister fixture class."""
    suite = unittest.TestSuite(loader.loadTestsFromTestCase(FbTemplateTests))
    for name in sorted(n for n in vars(ListerFbTests) if n.startswith('test_')):
        suite.addTest(ListerFbTests(name))
    return suite


if __name__ == '__main__':
    unittest.main()
